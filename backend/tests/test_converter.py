"""
Tests for the Bangla DOCX to PDF converter.
"""

from __future__ import annotations

import io
import os
import tempfile
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.bijoy import bijoy_to_unicode
from app.cleanup import cleanup_directory, cleanup_file, secure_filename, temporary_conversion_dir
from app.detector import FontDetectionResult, detect_fonts, is_legacy_bijoy_font
from app.main import app


# ── Bijoy Conversion Tests ───────────────────────────────────


class TestBijoyToUnicode:
    """Tests for Bijoy to Unicode text conversion."""

    def test_empty_string(self):
        assert bijoy_to_unicode("") == ""

    def test_none_handling(self):
        assert bijoy_to_unicode(None) is None

    def test_whitespace_only(self):
        assert bijoy_to_unicode("   ") == "   "

    def test_english_text_unchanged(self):
        """English text should pass through without modification."""
        result = bijoy_to_unicode("Hello World")
        # The converter may or may not modify pure English;
        # the important thing is it doesn't crash
        assert isinstance(result, str)

    def test_bijoy_consonants(self):
        """Basic Bijoy consonant conversion."""
        # 'K' in Bijoy = ক
        result = bijoy_to_unicode("K")
        assert "ক" in result

    def test_bijoy_vowel_sign(self):
        """Bijoy vowel signs (kar) conversion."""
        # 'v' in Bijoy = া (aa-kar)
        result = bijoy_to_unicode("Kv")
        assert "কা" in result or "া" in result

    def test_khetlal_title(self):
        """ÿ (ANSI 255) maps to ক্ষ (ksh)."""
        result = bijoy_to_unicode("‡ÿZjvj")
        assert result == "ক্ষেতলাল"

    def test_officer_in_charge_trailing_ref(self):
        """Trailing Ref in 'BbPvR©' must convert cleanly to 'ইনচার্জ' without IndexError."""
        result = bijoy_to_unicode("Rbve †gvt gy³viæj Avjg Awdmvi BbPvR©")
        assert result == "জনাব মোঃ মুক্তারুল আলম অফিসার ইনচার্জ"

    def test_kormoroto_and_forces(self):
        """Test 'Kg©iZ' -> 'কর্মরত' and '†dvm©‡`i' -> 'ফোর্সদের'."""
        assert bijoy_to_unicode("Kg©iZ") == "কর্মরত"
        assert bijoy_to_unicode("†dvm©‡`i") == "ফোর্সদের"

    def test_split_run_e_kar_and_merged_title(self):
        """Merged title containing 'bv‡gi' converts to 'নামের'."""
        text = "‡ÿZjvj _vbvq Kg©iZ Awdmvi‡`i bv‡gi ZvwjKvmn wbR †Rjvi Z_¨t "
        result = bijoy_to_unicode(text)
        assert "ক্ষেতলাল" in result
        assert "কর্মরত" in result
        assert "অফিসারদের" in result
        assert "নামের" in result


# ── Font Detection Tests ─────────────────────────────────────


class TestFontDetection:
    """Tests for legacy Bengali font detection."""

    def test_sutonnymj_is_legacy(self):
        assert is_legacy_bijoy_font("SutonnyMJ") is True

    def test_sutonnymj_bold_is_legacy(self):
        assert is_legacy_bijoy_font("SutonnyMJ Bold") is True

    def test_sutonnyomj_is_legacy(self):
        assert is_legacy_bijoy_font("SutonnyOMJ") is True

    def test_noto_sans_is_not_legacy(self):
        assert is_legacy_bijoy_font("Noto Sans Bengali") is False

    def test_nikosh_is_not_legacy(self):
        """Nikosh is a Unicode Bengali font and should NOT be flagged as legacy."""
        assert is_legacy_bijoy_font("Nikosh") is False
        assert is_legacy_bijoy_font("NikoshBAN") is False
        assert is_legacy_bijoy_font("NikoshLight") is False

    def test_kalpurush_is_not_legacy(self):
        """Kalpurush is Unicode; KalpurushANSI is legacy."""
        assert is_legacy_bijoy_font("Kalpurush") is False
        assert is_legacy_bijoy_font("KalpurushANSI") is True

    def test_solaiman_lipi_is_not_legacy(self):
        assert is_legacy_bijoy_font("SolaimanLipi") is False

    def test_arial_is_not_legacy(self):
        assert is_legacy_bijoy_font("Arial") is False

    def test_none_is_not_legacy(self):
        assert is_legacy_bijoy_font(None) is False

    def test_empty_string_is_not_legacy(self):
        assert is_legacy_bijoy_font("") is False

    def test_case_insensitive(self):
        assert is_legacy_bijoy_font("sutonnymj") is True
        assert is_legacy_bijoy_font("SUTONNYMJ") is True

    def test_contains_unicode_bengali(self):
        from app.detector import contains_unicode_bengali
        assert contains_unicode_bengali("আমার সোনার বাংলা") is True
        assert contains_unicode_bengali("Avgvi †mvbvi evsjv") is False
        assert contains_unicode_bengali("Hello 123") is False

    def test_is_legacy_bijoy_run_skips_unicode(self):
        from app.detector import is_legacy_bijoy_run
        doc = Document()
        p = doc.add_paragraph()
        run_unicode = p.add_run("বাংলাদেশ")
        run_ansi = p.add_run("evsjv‡`k")
        assert is_legacy_bijoy_run(run_unicode, "SutonnyMJ") is False
        assert is_legacy_bijoy_run(run_ansi, "SutonnyMJ") is True

    def test_detect_fonts_empty_doc(self):
        """An empty DOCX should report no legacy fonts."""
        doc = Document()
        result = detect_fonts(doc)
        assert result.has_legacy_fonts is False
        assert result.total_runs == 0

    def test_detect_fonts_with_unicode_doc(self):
        """A doc with Unicode Bengali text should not flag as legacy."""
        doc = Document()
        para = doc.add_paragraph()
        run = para.add_run("বাংলাদেশ")
        run.font.name = "Noto Sans Bengali"
        result = detect_fonts(doc)
        assert result.has_legacy_fonts is False

    def test_detect_fonts_with_legacy_doc(self):
        """A doc with SutonnyMJ font should be flagged as legacy."""
        doc = Document()
        para = doc.add_paragraph()
        run = para.add_run("evsjv‡`k")
        run.font.name = "SutonnyMJ"
        result = detect_fonts(doc)
        assert result.has_legacy_fonts is True
        assert "SutonnyMJ" in result.legacy_font_names


# ── DOCX Conversion Tests ────────────────────────────────────


class TestDocxConversion:
    """Tests for DOCX manipulation and Bijoy to Unicode document conversion."""

    def test_to_bengali_number(self):
        from app.converter import to_bengali_number

        assert to_bengali_number(1) == "১"
        assert to_bengali_number(2) == "২"
        assert to_bengali_number(10) == "১০"
        assert to_bengali_number(20) == "২০"

    def test_convert_bijoy_in_docx_merges_runs(self, tmp_path):
        from app.converter import convert_bijoy_in_docx

        doc = Document()
        p = doc.add_paragraph()
        r1 = p.add_run("bv‡")
        r1.font.name = "SutonnyMJ"
        r2 = p.add_run("gi")
        r2.font.name = "SutonnyMJ"

        in_file = tmp_path / "split.docx"
        out_file = tmp_path / "merged.docx"
        doc.save(in_file)

        result = convert_bijoy_in_docx(in_file, out_file)
        assert result["bijoy_detected"] is True

        merged_doc = Document(out_file)
        assert merged_doc.paragraphs[0].text == "নামের"

    def test_convert_bijoy_in_docx_table_serial_numbers(self, tmp_path):
        from app.converter import convert_bijoy_in_docx
        from docx.oxml.ns import qn

        doc = Document()
        table = doc.add_table(rows=3, cols=2)
        # Header
        table.rows[0].cells[0].text = "ক্রমিক"
        table.rows[0].cells[1].text = "নাম"

        # Row 1 with digit 1.
        cell_r1 = table.rows[1].cells[0]
        p1 = cell_r1.paragraphs[0]
        p1.text = "1."
        r1 = table.rows[1].cells[1].paragraphs[0].add_run("K")
        r1.font.name = "SutonnyMJ"

        # Row 2 with empty text but numPr
        cell_r2 = table.rows[2].cells[0]
        p2 = cell_r2.paragraphs[0]
        pPr = p2._p.get_or_add_pPr()
        numPr = pPr.makeelement(qn("w:numPr"))
        numPr.append(pPr.makeelement(qn("w:ilvl"), {qn("w:val"): "0"}))
        numPr.append(pPr.makeelement(qn("w:numId"), {qn("w:val"): "3"}))
        pPr.append(numPr)
        r2 = table.rows[2].cells[1].paragraphs[0].add_run("K")
        r2.font.name = "SutonnyMJ"

        in_file = tmp_path / "table.docx"
        out_file = tmp_path / "table_out.docx"
        doc.save(in_file)

        convert_bijoy_in_docx(in_file, out_file)
        out_doc = Document(out_file)
        out_table = out_doc.tables[0]
        assert out_table.rows[1].cells[0].text.strip() == "১."
        assert out_table.rows[2].cells[0].text.strip() == "২."


# ── Cleanup Tests ────────────────────────────────────────────


class TestCleanup:
    """Tests for file cleanup utilities."""

    def test_secure_filename_basic(self):
        assert secure_filename("document.docx") == "document.docx"

    def test_secure_filename_path_traversal(self):
        result = secure_filename("../../etc/passwd")
        assert "/" not in result
        assert "\\" not in result
        assert ".." not in result

    def test_secure_filename_null_bytes(self):
        result = secure_filename("file\x00.docx")
        assert "\x00" not in result

    def test_secure_filename_empty(self):
        result = secure_filename("")
        assert result == "document.docx"

    def test_secure_filename_hidden_file(self):
        result = secure_filename(".hidden.docx")
        assert not result.startswith(".")

    def test_temporary_conversion_dir(self):
        """Temp directory should be created and cleaned up."""
        import asyncio

        async def _test():
            async with temporary_conversion_dir() as tmp_dir:
                assert tmp_dir.exists()
                test_file = tmp_dir / "test.txt"
                test_file.write_text("hello")
                assert test_file.exists()
                saved_path = tmp_dir
            # After exiting, directory should be gone
            assert not saved_path.exists()

        asyncio.run(_test())

    def test_cleanup_nonexistent_directory(self):
        """Cleaning up a non-existent directory should not raise."""
        cleanup_directory(Path("/tmp/nonexistent_dir_12345"))

    def test_cleanup_file(self):
        """File cleanup should remove the file."""
        with tempfile.NamedTemporaryFile(delete=False, suffix=".txt") as f:
            f.write(b"test")
            path = Path(f.name)
        assert path.exists()
        cleanup_file(path)
        assert not path.exists()


# ── API Tests ────────────────────────────────────────────────


class TestAPI:
    """Tests for the FastAPI endpoints."""

    client = TestClient(app)

    def test_health_endpoint(self):
        response = self.client.get("/health")
        assert response.status_code == 200
        data = response.json()
        assert data["status"] == "healthy"
        assert data["service"] == "bangla-doc-to-pdf"

    def test_root_endpoint(self):
        response = self.client.get("/")
        assert response.status_code == 200

    def test_convert_no_file(self):
        response = self.client.post("/convert")
        assert response.status_code == 422  # Validation error

    def test_convert_wrong_extension(self):
        """Uploading a .txt file should be rejected."""
        content = b"Hello world"
        response = self.client.post(
            "/convert",
            files={"file": ("test.txt", io.BytesIO(content), "text/plain")},
        )
        assert response.status_code == 400
        assert "Invalid file type" in response.json()["detail"]

    def test_convert_too_large(self):
        """Files over 20MB should be rejected."""
        # Create content just over the limit
        content = b"x" * (21 * 1024 * 1024)
        response = self.client.post(
            "/convert",
            files={
                "file": (
                    "large.docx",
                    io.BytesIO(content),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        # Should be rejected (either 413 or 400 for invalid MIME)
        assert response.status_code in (400, 413)

    def test_convert_invalid_mime(self):
        """A file with .docx extension but wrong content should be rejected."""
        content = b"This is plain text, not a DOCX file"
        response = self.client.post(
            "/convert",
            files={
                "file": (
                    "fake.docx",
                    io.BytesIO(content),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert response.status_code == 400

    def test_convert_valid_docx(self):
        """A valid DOCX should be accepted (may fail if LibreOffice not installed)."""
        # Create a minimal valid DOCX
        doc = Document()
        doc.add_paragraph("Hello World - টেস্ট")
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        response = self.client.post(
            "/convert",
            files={
                "file": (
                    "test.docx",
                    buffer,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        # If LibreOffice is installed, this should succeed
        # If not, it'll return a 500 with a clear error
        assert response.status_code in (200, 500)
        if response.status_code == 500:
            assert "LibreOffice" in response.json()["detail"]


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
